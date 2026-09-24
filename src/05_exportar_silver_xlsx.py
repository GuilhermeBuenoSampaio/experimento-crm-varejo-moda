"""Exporta a Silver traduzida para Excel, preservando o Parquet original.

Uso: python src/05_exportar_silver_xlsx.py --silver-run-id SILVER01
Saida: data/03_exports/experimento_clientes_SILVER01.xlsx
"""
from __future__ import annotations

import argparse
import hashlib
import json
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
TITLES = {
    "hash_arquivo_fonte": "Hash do arquivo de origem",
    "numero_linha_fonte": "Numero da linha na origem",
    "meses_desde_ultima_compra": "Meses desde a ultima compra",
    "faixa_gasto_12_meses_usd": "Faixa de gasto em 12 meses (USD)",
    "gasto_12_meses_usd": "Gasto em 12 meses (USD)",
    "comprou_masculino_12_meses": "Comprou moda masculina em 12 meses",
    "comprou_feminino_12_meses": "Comprou moda feminina em 12 meses",
    "zona_localizacao": "Zona de localizacao",
    "cliente_novo_12_meses": "Cliente novo em 12 meses",
    "canal_compra_historico": "Canal historico de compra",
    "grupo_experimental": "Grupo experimental",
    "visitou_site_14_dias": "Visitou o site em 14 dias",
    "comprou_14_dias": "Comprou em 14 dias",
    "gasto_14_dias_usd": "Gasto em 14 dias (USD)",
}


def sha256(path: Path) -> str:
    value = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            value.update(chunk)
    return value.hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--silver-run-id", required=True)
    args = parser.parse_args()
    if not args.silver_run_id.isalnum():
        parser.error("ID da Silver deve conter apenas letras e numeros")

    try:
        import pyarrow.parquet as pq
        from openpyxl import Workbook
    except ImportError as error:
        raise SystemExit("Instale as dependencias: python -m pip install pyarrow openpyxl") from error

    audit_path = ROOT / "quality/runs" / args.silver_run_id / "validacao_silver.json"
    audit = json.loads(audit_path.read_text(encoding="utf-8"))
    if audit["validation"] != "approved" or audit["silver_run_id"] != args.silver_run_id:
        raise ValueError("Auditoria Silver nao aprovada ou ID divergente")
    parquet = ROOT / audit["parquet_relative_path"]
    if sha256(parquet) != audit["parquet_sha256"]:
        raise ValueError("Hash do Parquet diverge da auditoria")
    table = pq.read_table(parquet)
    if set(table.column_names) != set(TITLES):
        raise ValueError("Colunas inesperadas na Silver traduzida")
    if table.num_rows != audit["silver_records"]:
        raise ValueError("Contagem diverge da auditoria")
    if table.num_rows > 1048575:
        raise ValueError("Quantidade excede o limite de linhas de uma planilha Excel")

    wb = Workbook(write_only=True)
    data = wb.create_sheet("Base traduzida")
    data.append([TITLES[column] for column in table.column_names])
    columns = [column.to_pylist() for column in table.columns]
    for values in zip(*columns):
        data.append([float(value) if isinstance(value, Decimal) else value for value in values])

    dictionary = wb.create_sheet("Dicionario")
    dictionary.append(["Coluna na Silver", "Cabecalho no Excel"])
    for column in table.column_names:
        dictionary.append([column, TITLES[column]])

    metadata = wb.create_sheet("Rastreabilidade")
    for key, value in (
        ("Execucao Silver", args.silver_run_id),
        ("Linhas", table.num_rows),
        ("Arquivo Parquet", str(parquet.relative_to(ROOT))),
        ("SHA-256 Parquet", audit["parquet_sha256"]),
        ("Granularidade", "Um registro experimental por linha; nao existe ID de cliente."),
        ("Unidade monetaria", "USD"),
    ):
        metadata.append([key, value])

    destination = ROOT / "data/03_exports" / ("experimento_clientes_" + args.silver_run_id + ".xlsx")
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists():
        raise FileExistsError("Exportacao ja existe; nao sobrescrever: " + str(destination))
    wb.save(destination)
    print(json.dumps({"arquivo": str(destination), "linhas": table.num_rows,
                      "sha256_xlsx": sha256(destination)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
